"""Build the five-slide UHI presentation from the saved project results."""
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "artifacts" / "presentation"
OUT.mkdir(parents=True, exist_ok=True)
m = json.loads((ROOT / "data/app/metrics.json").read_text())
plans = json.loads((ROOT / "data/app/plans.json").read_text())
plan = plans["presets"]["10"]["ours"]
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BG, FG, MUTED, ORANGE, MINT, LINE = "080C0B", "F3F1E9", "AABAB1", "FF852F", "78D8B4", "34473D"


def rect(slide, x, y, w, h, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    return shape


def text(slide, x, y, w, h, value, size=18, color=FG, bold=False, font="Calibri"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = font
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(6)
    return box


def base(n, kicker, title, subtitle, source):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(BG)
    rect(s, .48, .38, .48, .045, ORANGE)
    text(s, 1.08, .28, 11.7, .28, kicker.upper(), 11, ORANGE, True, "Consolas")
    text(s, .5, .82, 12.3, .66, title, 32, bold=True)
    text(s, .52, 1.57, 12.1, .54, subtitle, 15, MUTED)
    rect(s, .5, 7.0, 12.33, .012, LINE)
    text(s, .5, 7.12, 11.5, .23, source, 9, MUTED)
    text(s, 12.13, 7.1, .7, .26, f"{n:02} / 05", 10, MINT, font="Consolas")
    return s


def section(s, x, y, w, title, body, size=17):
    text(s, x, y, w, .32, title, 15, MINT, True)
    text(s, x, y+.43, w, 1.2, body, size)


def image(s, path, x, y, w, h):
    from PIL import Image
    iw, ih = Image.open(path).size
    scale = min(w/iw, h/ih)
    s.shapes.add_picture(str(path), Inches(x+(w-iw*scale)/2), Inches(y+(h-ih*scale)/2),
                        width=Inches(iw*scale), height=Inches(ih*scale))


def notes(s, value):
    s.notes_slide.notes_text_frame.text = value


s = base(1, "HackMe’26 / AI–ML / Kochi", "UHI: Urban Heat Intelligence",
         "A decision-support dashboard that connects urban heat mapping to practical cooling decisions.",
         "Sources: README.md; PLAN.md; data/app/manifest.json; saved dashboard build.")
section(s, .55, 2.25, 4.35, "THE PROBLEM", "Heat is uneven across the city. Planners need to know which places and people face exposure, where money can help, and whether new development adds heat.", 19)
section(s, .55, 4.58, 4.35, "THE PROJECT’S ANSWER", "Locate hotspots → explain drivers → compare interventions → allocate budgets → screen development. Make the evidence visible at every step.", 19)
image(s, ROOT / "artifacts/m3/dashboard_desktop.png", 5.25, 2.2, 7.55, 4.12)
text(s, 5.3, 6.38, 7.3, .35, "74 named wards  /  54,168 study cells  /  23 satellite scenes", 13, MINT, font="Consolas")
notes(s, "UHI is built for municipal planners, ward representatives and project reviewers. Its five workspaces answer where heat is concentrated, what to change, what a scenario could do, whether development adds heat, and how reliable the evidence is. The wider study grid contains 54,168 cells. The 74 named wards are Kochi municipality, with 8,063 cells inside the municipal boundaries according to PLAN.md. Do not confuse municipal ward coverage with wider-grid optimization totals. Satellite values are land-surface temperature at approximately 10:30 AM during January–April, not the temperature people feel. Screenshot is from the local project build.")

s = base(2, "Data / models / decisions", "From satellite observations to an action plan",
         "A 100 m analytical grid combines surface conditions, weather, population and the built environment.",
         "Sources: README.md; src/uhi/model.py, scenarios.py, optimize.py; data/app/metrics.json.")
columns = [(.55, "01  OBSERVE", "Landsat 8/9 surface heat\nSentinel-2 + land cover\nERA5-Land atmosphere\nGHSL population + OSM"),
           (3.75, "02  EXPLAIN", "Scene-panel XGBoost\nPhysics-informed constraints\n(1 − albedo) × sunlight\nTreeSHAP heat drivers"),
           (6.95, "03  SIMULATE", "20-neighbour local analogs\nLabelled roof/pavement formulas\nJoint scenarios + spillover\nSupport checks and clipping"),
           (10.05, "04  PRIORITISE", "Person·°C benefit per rupee\nEligible public-land actions\n₹1 / ₹10 / ₹50 crore plans\nTypes, costs and locations")]
for x, title, body in columns:
    rect(s, x, 2.35, 2.72, .025, ORANGE)
    text(s, x, 2.55, 2.8, .37, title, 16, MINT, True, "Consolas")
    text(s, x, 3.13, 2.75, 1.75, body, 16)
rect(s, .55, 5.25, 12.23, .014, LINE)
section(s, .55, 5.5, 5.7, "EXPOSURE SCREENING", "Heat exposure combines surface-heat rank, weather-derived heat index and population; results roll up to wards.", 17)
section(s, 6.95, 5.5, 5.8, "WHY PHYSICS + ML?", "ML captures local patterns; explicit formulas cover interventions with weak analog evidence. Spatial validation tests transfer across areas.", 17)
notes(s, "Inputs are aligned to 100 m cells. The model includes per-scene atmospheric variables and a sunlight–albedo energy feature, with monotonic constraints. The CV artifact contains 250,000 scene-panel rows and 23 scenes. Analog interventions transition toward k=20 similar local cells, clip inputs to observed support and constrain changes toward cooling. Roof and pavement formulas are labelled separately and receive no neighbourhood spillover credit. The optimizer seeks total person-degrees of cooling per rupee, rather than simply the highest headcount. The dashboard reads committed frozen artifacts; moving a UI slider is a sensitivity preview, not a fresh full model training run. This distinction keeps the app responsive and its claims accurate.")

s = base(3, "Product / simulations / language", "One map. Six ways to explore a scenario.",
         "A map-first dashboard with ward selection, English/Malayalam interfaces and downloadable scenario reports.",
         "Sources: src/uhi/simulation_ui.py; simulation_preview.py; SIMULATIONS.md; browser-tested local UI.")
image(s, ROOT / "artifacts/m3/simulation_3d_desktop.png", .5, 2.16, 7.2, 4.6)
items = [
    ("INTERVENTIONS", "8 entries: trees, canal-bank strips, mangroves, ponds, cool/green roofs, cool pavements and zero-credit canals."),
    ("BUDGETS + DEVELOPMENT", "3 budgets with 2 baseline strategies; 6 project sites × 4 land uses, with available cooling offsets."),
    ("WARD LAB + EVIDENCE", "Adjust air temperature, humidity, wind and sunlight; inspect validation, data support and report outputs."),
    ("OPTIONAL 3D PREVIEW", "Start simulation, rotate the neighbourhood, replay changes and compare before/after. The scene is illustrative."),
]
for i, (title, body) in enumerate(items):
    y = 2.22+i*1.13
    text(s, 8.0, y, 4.8, .28, title, 13, MINT, True, "Consolas")
    text(s, 8.0, y+.34, 4.75, .79, body, 15)
notes(s, "The main dashboard has Heat overview, Cooling plan, Simulation, Project check and Evidence workspaces. Within Simulation, the six sections are Interventions, Budgets, Development, Ward lab, Evidence and 3D Preview. The latest optional 3D scene is a lightweight browser canvas that changes roofs, planting or water features according to the selected measure. Start simulation triggers the transition; later choices refresh it. Before/after, replay, a transition slider and drag rotation support demonstration. Geometry is schematic, not a surveyed digital twin; animation is not a forecast of time to maturity. Ward treatment numbers come from selected sites in the saved ₹50 crore plan where available; otherwise the study-wide median is labelled. Weather and treatment estimates are presented separately. Malayalam is standard written Malayalam.")

s = base(4, "Measured model performance / projected planning impact", "₹10 crore: a quantified cooling scenario",
         "Saved estimates for the wider study area, not measured post-construction outcomes.",
         "Sources: data/app/plans.json (preset 10); metrics.json (spatial_cv, backtest). Values rounded.")
stats = [(.55, f"{plan['people_cooled']:,.0f}", "people cooled ≥0.1°C"),
         (3.8, f"{abs(plan['mean_dt_cooled']):.2f}°C", "mean cooling in cooled cells"),
         (7.0, str(plan['cells']), "selected 100 m treatment sites"),
         (10.15, f"₹{plan['cost_rs']/1e7:.2f} cr", "estimated spend")]
for x, value, caption in stats:
    text(s, x, 2.22, 2.8, .68, value, 32, MINT, True)
    text(s, x, 2.95, 2.85, .6, caption, 14, MUTED)
text(s, .55, 3.82, 6.0, .4, "TOTAL COOLING BENEFIT / person·°C", 14, ORANGE, True, "Consolas")
strategies = [plan, *reversed(plans["presets"]["10"]["baselines"])]
maximum = plan["person_deg_cooling"]
for i, row in enumerate(strategies):
    y = 4.42+i*.66
    text(s, .55, y, 2.0, .3, row["strategy"], 15)
    rect(s, 2.65, y+.025, 3.0*row["person_deg_cooling"]/maximum, .23, MINT if i==0 else "52695D")
    text(s, 5.85, y-.015, 1.0, .35, f"{row['person_deg_cooling']:,.0f}", 15, MINT if i==0 else FG)
text(s, .55, 6.52, 6.25, .3, "Mix: 338 street-tree sites · 87 cool roofs · 5 mangrove · 1 canal strip", 12, MUTED)
rect(s, 7.2, 3.85, .015, 2.85, LINE)
text(s, 7.55, 3.83, 5.2, .35, "VALIDATION THAT CAN BE INSPECTED", 14, ORANGE, True, "Consolas")
text(s, 7.55, 4.35, 5.2, 1.28, "Spatial R² 0.829  |  RMSE 1.40°C\n140 held-out spatial blocks (~2 km)\nLinear baseline: R² 0.616; RMSE 2.10°C", 18)
text(s, 7.55, 5.85, 5.12, .85, "2017–2024 back-test: 8,000 changed cells; MAE 1.40°C; correlation 0.35. The unconstrained model scores slightly better (R² 0.839).", 15, MUTED)
notes(s, "Read this as a scenario estimate. The exact saved plan spend is ₹99,829,849.87, with 186,257.63 estimated people cooled by at least 0.1°C, mean cooling −1.20346°C and 228,650.34 person·°C. Person·°C weights cooling by the exposed population; it is not a count of unique interventions or a monetary benefit. The trees-everywhere baseline yields 44,297.60 person·°C and spread-evenly yields 7,108.48 at roughly the same budget. At ₹50 crore, trees everywhere reaches more people, while UHI has greater total person·°C benefit, so do not claim UHI maximizes every metric. The five mangrove sites have low analog support, which the UI flags. Physics-informed spatial R² is 0.82854 and RMSE 1.4032°C. Unconstrained XGBoost reaches R² 0.83911 and RMSE 1.35926°C: physics constraints are not claimed to improve raw accuracy. Historical back-testing is exploratory and does not establish intervention causality.")

s = base(5, "Trust / delivery / next steps", "Useful today. Explicit about what comes next.",
         "A working planning prototype with traceable outputs, local-language access and honest evidence boundaries.",
         "Sources: PLAN.md; SIMULATIONS.md; data/app/manifest.json; src/uhi/simulation.py; README.md.")
section(s, .55, 2.25, 5.75, "WHAT IS READY", "Interactive heat and ward views; saved cooling plans; 24 development scenarios; weather sensitivity; optional 3D preview; report downloads and standard Malayalam.", 19)
section(s, 7.0, 2.25, 5.7, "HOW TO INTERPRET IT", "Surface heat is not felt air temperature. City-scale weather is not a ward sensor. Scenarios are indicative; project screening is not planning permission.", 19)
section(s, .55, 4.34, 5.75, "DELIVERY + DEMO", "Python / Streamlit / XGBoost, serving a frozen data build. Demo: choose a ward → inspect heat → compare a plan → start 3D → download the report.", 18)
section(s, 7.0, 4.34, 5.7, "PRIORITY NEXT STEPS", "Obtain ECOSTRESS afternoon imagery and CPCB station data; produce the matched back-test; field-check costs and outcomes. SOLWEIG / InVEST remain future scope.", 18)
rect(s, .55, 6.48, 12.22, .025, ORANGE)
text(s, .55, 6.61, 12.1, .29, "THE VALUE: make heat decisions spatial, budget-aware, explainable and reviewable.", 16, MINT, True)
notes(s, "The deployed application can read committed data/app artifacts without Earth Engine at runtime. Rebuilding the real-data pipeline requires Earth Engine and source exports. The current manifest labels the source frozen and carries a build timestamp of 26 September 2026 at 02:11 IST. ECOSTRESS, CPCB and matched validation outputs are not present in this build and should never be presented as completed validation. Optional resident-reaction artifacts are also absent; SOLWEIG and InVEST are future work. Canal geometry is unverified OSM candidate data, not an official IURWTS alignment, and canals receive zero cooling credit. Before practical implementation, planners need site verification, permissions, locally checked costs, maintenance plans and field measurement. Suggested closing: UHI turns a heat map into an auditable conversation about where to act, what it may cost, what benefit is estimated and what remains uncertain.")

prs.core_properties.title = "UHI: Urban Heat Intelligence"
prs.core_properties.subject = "Five-slide project overview, simulations, results and evidence"
prs.core_properties.author = "UHI Project Team"
prs.core_properties.keywords = "Kochi, urban heat, geospatial AI, simulations, HackMe26"
target = OUT / "UHI_Project_Presentation.pptx"
prs.save(target)
check = Presentation(target)
assert len(check.slides) == 5
assert all(slide.notes_slide.notes_text_frame.text for slide in check.slides)
for slide in check.slides:
    for shape in slide.shapes:
        assert shape.left >= 0 and shape.top >= 0
        assert shape.left + shape.width <= prs.slide_width + 10
        assert shape.top + shape.height <= prs.slide_height + 10
print(target)
print(f"Verified {len(check.slides)} slides with speaker notes; {target.stat().st_size:,} bytes.")
